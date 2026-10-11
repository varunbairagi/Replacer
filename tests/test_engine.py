from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from landdocs.engine import (TEMPLATE, TEMPLATE_SHA256, TemplateError, cells,
                             generate_docx, land_tables, paragraph_edits,
                             replace_span, text, update_land_tables)
from landdocs.model import (ClientData, LandRecord, ValidationError,
                            load_client, save_client)

SAMPLE = Path(__file__).resolve().parent.parent / "sample_data/client.json"


class DocumentChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.client = load_client(SAMPLE)

    def generate(self, client=None):
        target = self.directory / "generated.docx"
        generate_docx(client or self.client, target)
        return target, Document(target)

    def test_original_and_all_other_package_members_unchanged(self):
        before = TEMPLATE.read_bytes()
        target, doc = self.generate()
        self.assertEqual(TEMPLATE.read_bytes(), before)
        self.assertEqual(sha256(before).hexdigest(), TEMPLATE_SHA256)
        with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(target) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != "word/document.xml":
                    self.assertEqual(a.read(name), b.read(name), name)
        source = Document(TEMPLATE)
        self.assertEqual(etree.tostring(source._element.body.sectPr), etree.tostring(doc._element.body.sectPr))

    def test_all_three_tables_grow_with_exact_strings_and_totals(self):
        records = tuple(LandRecord(f"000{i}/2/03", "0.10", "0.20") for i in range(7))
        target, doc = self.generate(replace(self.client, records=records))
        tables = list(land_tables(doc._element))
        self.assertEqual(len(tables), 3)
        self.assertTrue(any(a.tag == qn("w:tc") for a in tables[-1].iterancestors()))
        for table in tables:
            rows = table.findall(qn("w:tr"))
            self.assertEqual(len(rows), 9)
            for i, row in enumerate(rows[1:-1]):
                values = [text(c) for c in cells(row)]
                self.assertEqual(values[:4], [f"{i+1:02d}", records[i].khasra, "0.10", "0.20"])
                if len(values) == 5:
                    self.assertEqual(values[4], self.client.bahi if i == 0 else "")
                    self.assertEqual(cells(row)[4].xpath("./w:tcPr/w:vMerge")[0].get(qn("w:val")), "restart" if i == 0 else "continue")
            self.assertEqual([text(c) for c in cells(rows[-1])][:4], ["", "07 fdrk", "0.70", "1.40"])

    def test_single_record_shrinks_every_table_and_keeps_common_bahi_merge(self):
        _, doc = self.generate(replace(self.client, records=self.client.records[:1]))
        for table in land_tables(doc._element):
            rows = table.findall(qn("w:tr"))
            self.assertEqual(len(rows), 3)
            if len(cells(rows[0])) == 5:
                merge = cells(rows[1])[4].xpath("./w:tcPr/w:vMerge")[0]
                self.assertEqual(merge.get(qn("w:val")), "restart")
                self.assertTrue(cells(rows[-1])[4].xpath("./w:tcPr/w:vMerge"))

    def test_identity_residence_location_and_date_scope(self):
        client = replace(self.client, client_name="CLIENT", father_name="FATHER", address="ADDRESS", village="VILLAGE", tehsil="TEHSIL", district="DISTRICT", document_date="2028-02-29", bahi="0007/B", records=(LandRecord("003/1", "2.01", "0.1"),))
        _, doc = self.generate(client)
        content = text(doc._element)
        self.assertNotIn("o`Unkou", content)
        self.assertNotIn("txUukFk", content)
        self.assertNotIn("eaM[ksMk", content)
        self.assertNotIn("LD-445692", content)
        self.assertNotIn("1.229", content)
        self.assertEqual(content.count("CLIENT"), 13)
        self.assertEqual(content.count("FATHER"), 13)
        self.assertEqual(content.count("29.02.2028"), 6)
        self.assertIn("25@08@2026", content)
        self.assertIn("11-10-2022", content)
        self.assertIn("9827364038", content)
        self.assertIn("lhgksj@nksjkgk ftyk lhgksj", content)
        self.assertIn("'kk[kk vgeniqj ftyk lhgksj us CLIENT", content)
        self.assertIn("fuoklh ADDRESS rg0 TEHSIL ftyk DISTRICT", content)
        self.assertIn("fuoklh xzke ADDRESS o d`\"kd xzke VILLAGE", content)
        self.assertIn("esa 2.01 gSDV;j", content)

    def test_headers_grids_and_unedited_paragraphs_remain_identical(self):
        source = Document(TEMPLATE)
        _, target = self.generate()
        for a, b in zip(land_tables(source._element), land_tables(target._element)):
            for tag in ("w:tblPr", "w:tblGrid", "w:tr"):
                self.assertEqual(etree.tostring(a.find(qn(tag))), etree.tostring(b.find(qn(tag))))
        # All top-level paragraphs with no edits must retain exact XML.
        for a, b in zip(source.paragraphs, target.paragraphs):
            clone = deepcopy(a._p)
            if not paragraph_edits(clone, self.client):
                self.assertEqual(etree.tostring(a._p), etree.tostring(b._p))

    def test_total_suffix_keeps_hindi_font_and_dates_have_literal_separators(self):
        _, target = self.generate()
        source = Document(TEMPLATE)
        for a, b in zip(land_tables(source._element), land_tables(target._element)):
            original_cell = cells(a.findall(qn("w:tr"))[-1])[1]
            generated_cell = cells(b.findall(qn("w:tr"))[-1])[1]
            for old, new in zip(original_cell.iter(qn("w:r")), generated_cell.iter(qn("w:r"))):
                self.assertEqual(etree.tostring(old.find(qn("w:rPr"))), etree.tostring(new.find(qn("w:rPr"))))
            for cell in cells(b.findall(qn("w:tr"))[-1])[2:4]:
                self.assertFalse(cell.xpath(".//w:fldChar | .//w:instrText"))
        dates = [r for r in target._element.iter(qn("w:r")) if "22.09.2026" in text(r)]
        self.assertEqual(len(dates), 6)
        for run in dates:
            self.assertEqual(run.find("./"+qn("w:rPr")+"/"+qn("w:rFonts")).get(qn("w:ascii")), "Times New Roman")

    def test_unicode_is_explicit_and_has_no_devanagari_in_devlys_runs(self):
        client = replace(self.client, client_name="राम", father_name="मोहन", address="घर", village="ग्राम", tehsil="तहसील", district="जिला", text_mode="unicode", unicode_font="Mangal", caste="ब्राह्मण", bahi="००१", records=(LandRecord("००२/१", "0.1", "0.2"),))
        _, doc = self.generate(client)
        for run in doc._element.iter(qn("w:r")):
            if any(0x0900 <= ord(c) <= 0x097F for c in text(run)):
                self.assertEqual(run.find("./"+qn("w:rPr")+"/"+qn("w:rFonts")).get(qn("w:ascii")), client.unicode_font)
        self.assertIn("राम", text(doc._element))
        self.assertTrue(doc._element.xpath('.//w:rFonts[@w:ascii="DevLys 010"]'))

    def test_age_caste_and_document_bahi_cover_all_inspected_slots(self):
        client = replace(self.client, age="67", caste="CASTE", bahi="0007/B")
        _, doc = self.generate(client)
        content = text(doc._element)
        self.assertEqual(content.count("vk;q 67 tkfr CASTE"), 2)
        self.assertEqual(content.count("CASTE"), 10)
        self.assertNotIn("czkEg.k", content)
        self.assertNotIn("LD-445692", content)
        self.assertEqual(content.count("0007/B"), 5)  # Four prose slots plus merged table.
        for run in doc._element.iter(qn("w:r")):
            if "vk;q 67 tkfr" in text(run):
                self.assertEqual(run.find("./"+qn("w:rPr")+"/"+qn("w:rFonts")).get(qn("w:ascii")), "DevLys 010")
        raw = client.to_dict()
        self.assertEqual(raw["schema_version"], 2)
        self.assertEqual(raw["bahi"], "0007/B")
        self.assertTrue(all("bahi" not in r for r in raw["records"]))

    def test_old_json_migrates_shared_bahi_without_inventing_age_or_caste(self):
        raw = self.client.to_dict()
        raw["schema_version"] = 1
        bahi = raw.pop("bahi")
        raw.pop("age"); raw.pop("caste")
        for record in raw["records"]:
            record["bahi"] = bahi
        old = ClientData.from_dict(raw)
        self.assertEqual(old.bahi, bahi)
        self.assertEqual((old.age, old.caste), ("", ""))
        self.assertEqual(old.records, self.client.records)
        with self.assertRaises(ValidationError):
            generate_docx(old, self.directory / "incomplete.docx")
        self.assertFalse((self.directory / "incomplete.docx").exists())
        with self.assertRaises(ValidationError):
            save_client(old, self.directory / "incomplete.json")
        raw["records"][1]["bahi"] = "DIFFERENT"
        with self.assertRaisesRegex(ValidationError, "different Bahi"):
            ClientData.from_dict(raw)

    def test_age_caste_font_and_bahi_validation(self):
        for age in ("0", "151", "40.5", "-1", "NaN", "४०", 40):
            with self.subTest(age=age), self.assertRaises(ValidationError):
                replace(self.client, age=age)
        with self.assertRaises(ValidationError):
            replace(self.client, bahi="")
        with self.assertRaises(ValidationError):
            replace(self.client, caste="ब्राह्मण")
        with self.assertRaises(ValidationError):
            replace(self.client, unicode_font="Kohinoor Devanagari", text_mode="unicode")

    def test_run_split_replacement_keeps_prefix_suffix_properties(self):
        doc = Document()
        p = doc.add_paragraph()
        a = p.add_run("prefix ABC"); a.bold = True
        b = p.add_run("DEF suffix"); b.italic = True
        props = [etree.tostring(r._r.rPr) for r in p.runs]
        replace_span(p._p, 7, 13, "VALUE")
        self.assertEqual(p.text, "prefix VALUE suffix")
        self.assertEqual([etree.tostring(r._r.rPr) for r in p.runs], props)

    def test_unicode_in_same_run_isolates_font_and_preserves_suffix(self):
        p = Document().add_paragraph()
        p.add_run("before OLD after").font.name = "DevLys 010"
        replace_span(p._p, 7, 10, "राम", "Mangal")
        self.assertEqual(p.text, "before राम after")
        self.assertEqual([r.font.name for r in p.runs], ["DevLys 010", "Mangal", "DevLys 010"])

    def test_no_replacement_cascade_when_value_contains_template_tokens(self):
        client = replace(self.client, client_name="txUukFk eaM[ksMk")
        _, doc = self.generate(client)
        self.assertEqual(text(doc._element).count(client.client_name), 13)

    def test_deep_nested_land_table_traversal(self):
        doc = Document(TEMPLATE)
        table = deepcopy(list(land_tables(doc._element))[0])
        container = Document()
        outer = container.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0)._tc.append(table)
        self.assertEqual(update_land_tables(container._element, replace(self.client, records=self.client.records[:2])), 1)
        self.assertEqual(len(table.findall(qn("w:tr"))), 4)

    def test_refuses_source_existing_and_uninspected_templates(self):
        before = TEMPLATE.read_bytes()
        with self.assertRaises(ValidationError):
            generate_docx(self.client, TEMPLATE)
        existing = self.directory / "exists.docx"
        existing.write_bytes(b"keep me")
        with self.assertRaises(FileExistsError):
            generate_docx(self.client, existing)
        self.assertEqual(existing.read_bytes(), b"keep me")
        altered = self.directory / "altered.docx"
        altered.write_bytes(before + b"extra")
        with self.assertRaises(TemplateError):
            generate_docx(self.client, self.directory / "new.docx", altered)
        self.assertFalse((self.directory / "new.docx").exists())
        self.assertEqual(TEMPLATE.read_bytes(), before)

    def test_decimal_and_json_round_trip_keep_identifiers_and_precision(self):
        client = replace(self.client, bahi="0001", records=(LandRecord("002/03", "0.10", "0.20"), LandRecord("002/04", "0.20", "0.10")))
        self.assertEqual(client.totals, (Decimal("0.30"), Decimal("0.30")))
        target = self.directory / "client.json"
        save_client(client, target)
        self.assertEqual(load_client(target), client)
        self.assertEqual(json.loads(target.read_text())["records"][0]["khasra"], "002/03")

    def test_invalid_data_is_rejected_before_output(self):
        for value in ("NaN", "Infinity", "-1", "1e2", "1,000", "0.123456789", 0.1, ""):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                LandRecord("002", value, "0")
        with self.assertRaises(ValidationError):
            replace(self.client, document_date="2025-02-29")
        with self.assertRaises(ValidationError):
            replace(self.client, records=())
        with self.assertRaises(ValidationError):
            replace(self.client, client_name="राम")
        with self.assertRaises(ValidationError):
            replace(self.client, records=(LandRecord("राम", "1", "1"),))
        with self.assertRaises(ValidationError):
            replace(self.client, text_mode="unicode", unicode_font="DevLys 010")
        raw = self.client.to_dict(); raw["records"][0]["khasra"] = 123
        with self.assertRaises(ValidationError):
            ClientData.from_dict(raw)
        raw = self.client.to_dict(); raw["unexpected"] = True
        with self.assertRaises(ValidationError):
            ClientData.from_dict(raw)


if __name__ == "__main__":
    unittest.main()
